import asyncio
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo

import pytest
from fastapi import HTTPException, Response
from openpyxl import load_workbook
from starlette.requests import Request

from app.demo_data import DEMO_MODEL, DEMO_MODEL_PATH
from app.main import analyze_model, compare_model_exports, export_excel
from app.observability import log_http_request
from app.rate_limit import InMemoryRateLimiter, RateLimitRule
from app.security import assert_safe_origin, cors_origins, docs_enabled, validate_security_config
from app.schemas import CompareResponse
from app.services.analyzer import PowerBIAnalyzer
from app.services.compare import compare_models
from app.services.excel_export import build_excel
from app.ingestion.loader import read_model_upload
from app.ingestion import pbip_reader
from app.ingestion.pbip_reader import read_pbip_archive
from app.upload_validation import max_pbip_upload_bytes, read_json_upload, validate_model_export


class FakeUpload:
    def __init__(self, filename: str, content: bytes):
        self.filename = filename
        self._content = content
        self._position = 0

    async def read(self, size: int = -1) -> bytes:
        if size < 0:
            result = self._content[self._position:]
            self._position = len(self._content)
            return result
        result = self._content[self._position:self._position + size]
        self._position += len(result)
        return result


PBIP_FIXTURE_ROOT = Path(__file__).parent / "fixtures" / "pbip_tmsl"
TMDL_FIXTURE_ROOT = Path(__file__).parent / "fixtures" / "pbip_tmdl"
JSON_EQUIVALENT_PATH = Path(__file__).parent / "fixtures" / "json_equivalent_model.json"


def make_pbip_zip(*extra_entries: tuple[str, bytes]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        for path in PBIP_FIXTURE_ROOT.rglob("*"):
            if path.is_file():
                archive.writestr(path.relative_to(PBIP_FIXTURE_ROOT).as_posix(), path.read_bytes())
        for path, content in extra_entries:
            archive.writestr(path, content)
    return output.getvalue()


def make_semantic_model_zip() -> bytes:
    output = BytesIO()
    semantic_root = PBIP_FIXTURE_ROOT / "Demo.SemanticModel"
    with ZipFile(output, "w") as archive:
        for path in semantic_root.rglob("*"):
            if path.is_file():
                archive.writestr(path.relative_to(semantic_root).as_posix(), path.read_bytes())
    return output.getvalue()


def make_tmdl_zip(*extra_entries: tuple[str, bytes]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        for path in TMDL_FIXTURE_ROOT.rglob("*"):
            if path.is_file():
                archive.writestr(path.relative_to(TMDL_FIXTURE_ROOT).as_posix(), path.read_bytes())
        for path, content in extra_entries:
            archive.writestr(path, content)
    return output.getvalue()


def make_zip(entries: list[tuple[str, bytes]]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        for path, content in entries:
            archive.writestr(path, content)
    return output.getvalue()


def make_request(method: str = "GET", origin: str | None = None, referer: str | None = None, request_id: str | None = None) -> Request:
    headers = []
    if origin:
        headers.append((b"origin", origin.encode()))
    if referer:
        headers.append((b"referer", referer.encode()))
    if request_id:
        headers.append((b"x-request-id", request_id.encode()))
    return Request({"type": "http", "method": method, "path": "/", "headers": headers, "client": ("127.0.0.1", 50000)})


def make_path_request(path: str, method: str = "POST", client: str = "127.0.0.1") -> Request:
    return Request({"type": "http", "method": method, "path": path, "headers": [], "client": (client, 50000)})


def test_analyzer_demo_report_has_expected_sections():
    report = PowerBIAnalyzer(DEMO_MODEL).full_report()
    assert DEMO_MODEL_PATH.name == "public_demo_model.json"
    assert report["summary"]["Dashboard"] == "Demo Publica Comercial"
    assert report["summary"]["Tabelas totais"] == 5
    assert "SQL" in report["summary"]["Tipos de fontes"]
    assert "Excel" in report["summary"]["Tipos de fontes"]
    assert "SharePoint" in report["summary"]["Tipos de fontes"]
    assert report["sources"][0]["Servidor"] == "demo-sql-server"
    assert report["relationships"][-1]["Ativo"] == "Não"
    assert "quality" not in report


def test_analyzer_filters_technical_date_tables():
    data = deepcopy(DEMO_MODEL)
    data["tables"].append({"name": "LocalDateTable_123", "columns": [{"name": "Date"}]})
    assert all(row["Tabela"] != "LocalDateTable_123" for row in PowerBIAnalyzer(data).full_report()["tables"])


def test_compare_models_reports_measure_column_and_relationship_changes():
    base = deepcopy(DEMO_MODEL)
    new = deepcopy(DEMO_MODEL)
    new["dashboardName"] = "Demo Comercial v2"
    new["tables"][0]["measures"][0]["expression"] = "SUM('Fato Vendas Demo'[Receita]) * 1.1"
    new["tables"][0]["measures"].append({"name": "Receita Nova", "expression": "SUM('Fato Vendas Demo'[Receita Nova])"})
    new["tables"][0]["columns"][5]["formatString"] = "R$ #,0"
    new["relationships"][0]["crossFilteringBehavior"] = "BothDirections"
    result = compare_models(base, new)
    response = CompareResponse.model_validate(result)
    assert result["dashboard_novo"] == "Demo Comercial v2"
    assert response.medidas.adicionadas[0].expressao_dax == "SUM('Fato Vendas Demo'[Receita Nova])"
    assert result["medidas"]["modificadas"][0]["medida"] == "Receita Total"
    assert result["colunas"]["modificadas"][0]["coluna"] == "Receita"
    assert result["relacionamentos"]["modificados"][0]["relacionamento"] == "Fato Vendas Demo.Id Produto -> Dim Produto Demo.Id Produto"


def test_build_excel_includes_expected_sheets():
    workbook = load_workbook(build_excel(PowerBIAnalyzer(DEMO_MODEL).full_report()))
    assert "Resumo" in workbook.sheetnames
    assert "Qualidade" not in workbook.sheetnames
    assert "Colunas em Medidas" in workbook.sheetnames
    assert workbook["Resumo"]["A1"].value == "Indicador"


def test_validate_model_export_rejects_missing_tables():
    with pytest.raises(HTTPException) as exc:
        validate_model_export({"dashboardName": "Sem tabelas"})
    assert exc.value.status_code == 422
    assert "tables" in exc.value.detail


def test_read_json_upload_rejects_invalid_json():
    with pytest.raises(HTTPException) as exc:
        asyncio.run(read_json_upload(FakeUpload("modelo.json", b"{invalid")))
    assert exc.value.status_code == 400
    assert "JSON inválido" in exc.value.detail


def test_read_json_upload_rejects_non_json_empty_and_large_files(monkeypatch):
    for upload in (FakeUpload("modelo.txt", b"{}"), FakeUpload("modelo.json", b"")):
        with pytest.raises(HTTPException) as exc:
            asyncio.run(read_json_upload(upload))
        assert exc.value.status_code == 400

    monkeypatch.setenv("LEITORBI_MAX_UPLOAD_MB", "1")
    with pytest.raises(HTTPException) as exc:
        asyncio.run(read_json_upload(FakeUpload("modelo.json", b" " * (1024 * 1024 + 1))))
    assert exc.value.status_code == 413


def test_read_pbip_upload_normalizes_tmsl_model_bim():
    payload = make_pbip_zip()
    data = asyncio.run(read_model_upload(FakeUpload("Demo.zip", payload)))
    report = PowerBIAnalyzer(data).full_report()

    assert report["raw"]["dashboardName"] == "Demo"
    assert report["raw"]["modelName"] == "Modelo TMSL Demo"
    assert report["summary"]["Cultura"] == "pt-BR"
    assert report["summary"]["Modo padrao"] == "import"
    assert report["measures"][0]["Expressao DAX"] == "SUM('Vendas'[Receita])"
    assert report["sources"][0]["Servidor"] == "demo-server"
    assert report["relationships"][0]["Relacionamento"] == "Vendas Calendario"


def test_read_pbip_upload_accepts_semantic_model_only_zip():
    data = asyncio.run(read_model_upload(FakeUpload("semantic-model.zip", make_semantic_model_zip())))
    assert data["modelName"] == "Modelo TMSL Demo"
    assert len(data["tables"]) == 2


def test_read_model_upload_does_not_accept_pbix():
    with pytest.raises(HTTPException) as exc:
        asyncio.run(read_model_upload(FakeUpload("modelo.pbix", b"PBIX")))
    assert exc.value.status_code == 400
    assert "JSON" in exc.value.detail
    assert "ZIP" in exc.value.detail


def test_pbip_analysis_comparison_and_excel_use_existing_contracts():
    payload = make_pbip_zip()
    analyzed = asyncio.run(analyze_model(FakeUpload("Demo.zip", payload)))
    compared = asyncio.run(compare_model_exports(FakeUpload("base.zip", payload), FakeUpload("novo.zip", payload)))
    exported = asyncio.run(export_excel(FakeUpload("Demo.zip", payload)))

    assert analyzed.raw.modelName == "Modelo TMSL Demo"
    assert compared.dashboard_base == "Demo"
    assert compared.tabelas.adicionadas == []
    assert exported.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_pbip_and_json_produce_equivalent_logical_analysis():
    json_model = json.loads(JSON_EQUIVALENT_PATH.read_text(encoding="utf-8"))
    pbip_model = asyncio.run(read_model_upload(FakeUpload("Demo.zip", make_pbip_zip())))

    json_report = PowerBIAnalyzer(json_model).full_report()
    pbip_report = PowerBIAnalyzer(pbip_model).full_report()

    for key in ("summary", "tables", "columns", "measures", "sources", "relationships", "columnsUsedInMeasures"):
        assert pbip_report[key] == json_report[key]


def test_pbip_ignores_report_cache_and_other_non_model_artifacts():
    payload = make_pbip_zip(
        ("Demo.Report/definition.pbir", b"not analyzed"),
        ("Demo.SemanticModel/cache.abf", b"sensitive cache bytes"),
        ("DAXQueries/query.dax", b"not analyzed"),
        ("TMDLScripts/script.tmdl", b"not analyzed"),
    )
    data = asyncio.run(read_model_upload(FakeUpload("Demo.zip", payload)))
    assert data["modelName"] == "Modelo TMSL Demo"


def test_read_pbip_upload_normalizes_tmdl_model():
    data = asyncio.run(read_model_upload(FakeUpload("tmdl.zip", make_tmdl_zip())))
    report = PowerBIAnalyzer(data).full_report()

    assert report["raw"]["dashboardName"] == "Demo"
    assert report["raw"]["modelName"] == "Modelo TMSL Demo"
    assert report["summary"]["Cultura"] == "pt-BR"
    assert report["summary"]["Modo padrao"] == "import"
    assert report["measures"][0]["Expressao DAX"].startswith("VAR Total")
    assert "SUM('Vendas'[Receita])" in report["measures"][0]["Expressao DAX"]
    assert report["sources"][0]["Servidor"] == "demo-server"
    assert report["relationships"][0]["Ativo"] == "Sim"
    assert report["columnsUsedInMeasures"] == [{"Tabela": "Vendas", "Coluna": "Receita"}]


def test_tmdl_matches_tmsl_and_json_logical_analysis():
    json_model = json.loads(JSON_EQUIVALENT_PATH.read_text(encoding="utf-8"))
    tmsl_model = asyncio.run(read_model_upload(FakeUpload("tmsl.zip", make_pbip_zip())))
    tmdl_model = asyncio.run(read_model_upload(FakeUpload("tmdl.zip", make_tmdl_zip())))

    json_report = PowerBIAnalyzer(json_model).full_report()
    tmsl_report = PowerBIAnalyzer(tmsl_model).full_report()
    tmdl_report = PowerBIAnalyzer(tmdl_model).full_report()
    for key in ("summary", "tables", "columns", "relationships", "columnsUsedInMeasures"):
        assert tmdl_report[key] == tmsl_report[key]
    for tmdl_source, tmsl_source in zip(tmdl_report["sources"], tmsl_report["sources"]):
        assert {key: value for key, value in tmdl_source.items() if key != "Expressao M"} == {
            key: value for key, value in tmsl_source.items() if key != "Expressao M"
        }
        assert "".join(tmdl_source["Expressao M"].split()) == "".join(tmsl_source["Expressao M"].split())
    assert tmdl_report["measures"][0]["Medida"] == tmsl_report["measures"][0]["Medida"]
    assert "SUM('Vendas'[Receita])" in tmdl_report["measures"][0]["Expressao DAX"]
    assert {key: value for key, value in tmdl_report["measures"][0].items() if key not in {"Expressao DAX", "Tamanho DAX"}} == {
        key: value for key, value in tmsl_report["measures"][0].items() if key not in {"Expressao DAX", "Tamanho DAX"}
    }
    assert json_report["tables"] == tmsl_report["tables"]


def test_tmdl_comparison_and_excel_use_existing_contracts():
    payload = make_tmdl_zip()
    analyzed = asyncio.run(analyze_model(FakeUpload("tmdl.zip", payload)))
    compared = asyncio.run(compare_model_exports(FakeUpload("base.zip", payload), FakeUpload("novo.zip", payload)))
    exported = asyncio.run(export_excel(FakeUpload("tmdl.zip", payload)))

    assert analyzed.raw.modelName == "Modelo TMSL Demo"
    assert compared.tabelas.adicionadas == []
    assert exported.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_tmdl_can_be_compared_with_json_through_existing_compare_models():
    json_payload = JSON_EQUIVALENT_PATH.read_bytes()
    compared = asyncio.run(
        compare_model_exports(
            FakeUpload("base.json", json_payload),
            FakeUpload("novo-tmdl.zip", make_tmdl_zip()),
        )
    )
    assert compared.tabelas.adicionadas == []
    assert compared.colunas.adicionadas == []


def test_tmdl_semantic_model_only_zip_is_supported():
    output = BytesIO()
    semantic_root = TMDL_FIXTURE_ROOT / "Demo.SemanticModel"
    with ZipFile(output, "w") as archive:
        for path in semantic_root.rglob("*"):
            if path.is_file():
                archive.writestr(path.relative_to(semantic_root).as_posix(), path.read_bytes())
    data = asyncio.run(read_model_upload(FakeUpload("semantic-tmdl.zip", output.getvalue())))
    assert data["modelName"] == "Modelo TMSL Demo"
    assert len(data["tables"]) == 2


@pytest.mark.parametrize(
    "entries, expected",
    [
        ([
            ("Demo.SemanticModel/definition/model.tmdl", b"model Demo\n    culture: pt-BR\n"),
        ], "nenhuma tabela"),
        ([
            ("Demo.SemanticModel/definition/model.tmdl", b"model Demo\n    culture: pt-BR\n"),
            ("Demo.SemanticModel/definition/tables/Broken.tmdl", b"table 'Broken'\n    measure 'M' = ```\n        SUM(1)\n"),
        ], "bloco multilinha"),
        ([
            ("Demo.SemanticModel/definition/model.tmdl", b"model Demo\n"),
            ("Demo.SemanticModel/definition/tables/Broken.tmdl", b"table\n"),
        ], "entidade sem nome"),
        ([
            ("Demo.SemanticModel/definition/model.tmdl", b"model Demo\n"),
            ("Demo.SemanticModel/definition/tables/Broken.tmdl", b"table 'Broken'\n    column 'Id'\n        dataType: int64\n"),
            ("Demo.SemanticModel/definition/relationships.tmdl", b"relationship 'Broken'\n    fromColumn: 'Missing'.'Id'\n    toColumn: 'Broken'.'Id'\n"),
        ], "tabela inexistente"),
    ],
)
def test_tmdl_invalid_structures_are_rejected(entries, expected):
    with pytest.raises(HTTPException) as exc:
        read_pbip_archive(make_zip(entries), "invalid-tmdl.zip")
    assert exc.value.status_code == 422
    assert expected in exc.value.detail


def test_pbip_rejects_conflicting_tmsl_and_tmdl_structures():
    with pytest.raises(HTTPException) as exc:
        read_pbip_archive(
            make_pbip_zip(("Demo.SemanticModel/definition/model.tmdl", b"model Demo")),
            "conflicting.zip",
        )
    assert exc.value.status_code == 422
    assert "TMSL e TMDL" in exc.value.detail


def test_read_pbip_archive_rejects_invalid_structure_and_unsafe_paths():
    missing_archive = BytesIO()
    with ZipFile(missing_archive, "w") as archive:
        archive.writestr("Other/file.txt", b"x")
    with pytest.raises(HTTPException) as missing:
        asyncio.run(read_model_upload(FakeUpload("empty.zip", missing_archive.getvalue())))
    assert missing.value.status_code == 422
    assert "model.bim" in missing.value.detail

    unsafe = BytesIO()
    with ZipFile(unsafe, "w") as archive:
        archive.writestr(ZipInfo("../escape.txt"), b"x")
    with pytest.raises(HTTPException) as traversal:
        read_pbip_archive(unsafe.getvalue(), "unsafe.zip")
    assert traversal.value.status_code == 422
    assert "inseguro" in traversal.value.detail


def test_read_pbip_archive_rejects_nested_archives():
    nested = BytesIO()
    with ZipFile(nested, "w") as archive:
        archive.writestr("inner.zip", b"PK")
    with pytest.raises(HTTPException) as exc:
        read_pbip_archive(nested.getvalue(), "nested.zip")
    assert exc.value.status_code == 422
    assert "compactados aninhados" in exc.value.detail


def test_read_pbip_archive_reports_live_connection_without_local_model():
    report_only = BytesIO()
    with ZipFile(report_only, "w") as archive:
        archive.writestr("Demo.Report/definition.pbir", b"{}")
    with pytest.raises(HTTPException) as exc:
        read_pbip_archive(report_only.getvalue(), "Demo.zip")
    assert "modelo semântico local" in exc.value.detail


def test_read_pbip_archive_sanitizes_malicious_upload_filename():
    data = read_pbip_archive(make_pbip_zip(), "..\\secret\\malicious.zip")
    assert data["dashboardName"] == "Demo"


def test_read_pbip_archive_rejects_corrupt_zip():
    with pytest.raises(HTTPException) as exc:
        read_pbip_archive(b"not-a-zip", "corrupt.zip")
    assert exc.value.status_code == 422
    assert "ZIP inválido ou corrompido" in exc.value.detail


def test_read_pbip_archive_reports_protected_model(monkeypatch):
    payload = make_pbip_zip()
    original_read = ZipFile.read

    def protected_read(self, *args, **kwargs):
        if args and getattr(args[0], "filename", "").endswith("model.bim"):
            raise RuntimeError("password required")
        return original_read(self, *args, **kwargs)

    monkeypatch.setattr(ZipFile, "read", protected_read)
    with pytest.raises(HTTPException) as exc:
        read_pbip_archive(payload, "protected.zip")
    assert exc.value.status_code == 422
    assert "protegido" in exc.value.detail


def test_read_pbip_archive_enforces_file_and_uncompressed_limits(monkeypatch):
    too_many = BytesIO()
    with ZipFile(too_many, "w") as archive:
        archive.writestr("one.txt", b"1")
        archive.writestr("two.txt", b"2")

    monkeypatch.setattr(pbip_reader, "MAX_PBIP_FILES", 1)
    with pytest.raises(HTTPException) as file_limit:
        read_pbip_archive(too_many.getvalue(), "too-many.zip")
    assert "quantidade de arquivos" in file_limit.value.detail

    expanded = BytesIO()
    with ZipFile(expanded, "w") as archive:
        archive.writestr("payload.txt", b"12")

    monkeypatch.setattr(pbip_reader, "MAX_PBIP_FILES", 5000)
    monkeypatch.setattr(pbip_reader, "MAX_PBIP_UNCOMPRESSED_BYTES", 1)
    with pytest.raises(HTTPException) as size_limit:
        read_pbip_archive(expanded.getvalue(), "expanded.zip")
    assert "descompactado" in size_limit.value.detail


def test_read_pbip_upload_enforces_separate_upload_limit(monkeypatch):
    monkeypatch.setenv("LEITORBI_MAX_PBIP_UPLOAD_MB", "1")
    with pytest.raises(HTTPException) as exc:
        asyncio.run(read_model_upload(FakeUpload("large.zip", b"x" * (1024 * 1024 + 1))))
    assert exc.value.status_code == 413
    assert "1 MB" in exc.value.detail


@pytest.mark.parametrize("size_mb", [1, 25, 99])
def test_pbip_upload_sizes_below_100_mb_are_received_incrementally(monkeypatch, size_mb):
    monkeypatch.setenv("LEITORBI_MAX_PBIP_UPLOAD_MB", "100")
    monkeypatch.setattr(
        "app.upload_validation.read_pbip_archive",
        lambda _content, _filename: {"tables": [{"name": "Teste"}], "relationships": []},
    )
    result = asyncio.run(read_model_upload(FakeUpload("padded.zip", b"x" * (size_mb * 1024 * 1024))))
    assert result["tables"][0]["name"] == "Teste"


def test_pbip_upload_above_100_mb_is_rejected_before_zip_read(monkeypatch):
    monkeypatch.setenv("LEITORBI_MAX_PBIP_UPLOAD_MB", "100")
    monkeypatch.setattr(
        "app.upload_validation.read_pbip_archive",
        lambda *_args: pytest.fail("ZIP não deveria ser lido após exceder o limite"),
    )
    with pytest.raises(HTTPException) as exc:
        asyncio.run(read_model_upload(FakeUpload("too-large.zip", b"x" * (100 * 1024 * 1024 + 1))))
    assert exc.value.status_code == 413


def test_default_pbip_upload_limit_is_100_mb(monkeypatch):
    monkeypatch.delenv("LEITORBI_MAX_PBIP_UPLOAD_MB", raising=False)
    assert max_pbip_upload_bytes() == 100 * 1024 * 1024


@pytest.mark.parametrize("table", [{}, {"name": "  "}])
def test_validate_model_export_rejects_table_without_name(table):
    with pytest.raises(HTTPException) as exc:
        validate_model_export({"tables": [table]})
    assert exc.value.status_code == 422


def test_request_observability_adds_request_id_and_logs(caplog):
    async def call_next(_request: Request) -> Response:
        return Response(status_code=204)

    request = make_request(request_id="teste-request-id")
    with caplog.at_level("INFO", logger="leitorbi.api"):
        response = asyncio.run(log_http_request(request, call_next))
    assert response.headers["X-Request-ID"] == "teste-request-id"
    records = [record for record in caplog.records if record.name == "leitorbi.api"]
    assert records[-1].message == "request_completed"
    assert records[-1].request_id == "teste-request-id"
    assert records[-1].status_code == 204


def test_rate_limiter_is_per_client_and_expires(monkeypatch):
    monkeypatch.delenv("LEITORBI_RATE_LIMIT_ANALYZE", raising=False)
    limiter = InMemoryRateLimiter({"/api/models/analyze": RateLimitRule(2, window_seconds=60)})
    request = make_path_request("/api/models/analyze")
    other_client = make_path_request("/api/models/analyze", client="192.0.2.10")

    assert limiter.check(request, now=0) == (True, 0)
    assert limiter.check(request, now=1) == (True, 0)
    allowed, retry_after = limiter.check(request, now=2)
    assert allowed is False
    assert 57 <= retry_after <= 60
    assert limiter.check(other_client, now=2) == (True, 0)
    assert limiter.check(request, now=61) == (True, 0)


def test_public_mutations_still_validate_origin(monkeypatch):
    monkeypatch.setenv("LEITORBI_CORS_ORIGINS", "http://localhost:5173")
    assert_safe_origin(make_request(method="POST", origin="http://localhost:5173"))
    with pytest.raises(HTTPException) as exc:
        assert_safe_origin(make_request(method="POST", origin="https://evil.example"))
    assert exc.value.status_code == 403


def test_origin_requirement_and_cors_configuration(monkeypatch):
    monkeypatch.setenv("LEITORBI_REQUIRE_ORIGIN", "true")
    with pytest.raises(HTTPException) as exc:
        assert_safe_origin(make_request(method="POST"))
    assert exc.value.status_code == 403

    monkeypatch.setenv("LEITORBI_CORS_ORIGINS", "*")
    assert cors_origins() == ["*"]
    monkeypatch.setenv("LEITORBI_CORS_ORIGINS", "*,http://localhost:5173")
    with pytest.raises(RuntimeError):
        cors_origins()


def test_production_disables_api_docs_by_default(monkeypatch):
    monkeypatch.setenv("LEITORBI_ENV", "production")
    monkeypatch.delenv("LEITORBI_ENABLE_DOCS", raising=False)
    assert docs_enabled() is False
    monkeypatch.setenv("LEITORBI_ENABLE_DOCS", "true")
    assert docs_enabled() is True


def test_production_security_requires_explicit_cors_but_not_session_settings(monkeypatch):
    monkeypatch.setenv("LEITORBI_ENV", "production")
    monkeypatch.delenv("LEITORBI_CORS_ORIGINS", raising=False)
    with pytest.raises(RuntimeError):
        validate_security_config()
    monkeypatch.setenv("LEITORBI_CORS_ORIGINS", "https://leitorbi.example")
    validate_security_config()


def test_model_operations_are_public_without_cookies():
    payload = json.dumps(DEMO_MODEL).encode()
    analyzed = asyncio.run(analyze_model(FakeUpload("modelo.json", payload)))
    compared = asyncio.run(compare_model_exports(FakeUpload("base.json", payload), FakeUpload("novo.json", payload)))
    exported = asyncio.run(export_excel(FakeUpload("modelo.json", payload)))
    assert analyzed.raw.dashboardName == "Demo Publica Comercial"
    assert compared.dashboard_base == "Demo Publica Comercial"
    assert exported.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
