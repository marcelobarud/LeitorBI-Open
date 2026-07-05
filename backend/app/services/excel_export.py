from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.utils import get_column_letter


def text(value: Any) -> str:
    return "" if value is None else str(value)


def build_excel(report: dict[str, Any], selected_sheet: str | None = None) -> BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheets = {
        "Resumo": [report.get("summary", {})],
        "Tabelas": report.get("tables", []),
        "Colunas": report.get("columns", []),
        "Medidas": report.get("measures", []),
        "Fontes": report.get("sources", []),
        "Relacionamentos": report.get("relationships", []),
        "Colunas em Medidas": report.get("columnsUsedInMeasures", []),
    }
    if selected_sheet:
        sheets = {selected_sheet: sheets.get(selected_sheet, [])}

    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_style = NamedStyle(name="leitorbi_header")
    header_style.font = Font(bold=True, color="FFFFFF")
    header_style.fill = PatternFill("solid", fgColor="1F4E79")
    header_style.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_style.border = border

    body_style = NamedStyle(name="leitorbi_body")
    body_style.alignment = Alignment(vertical="top", wrap_text=True)
    body_style.border = border

    workbook.add_named_style(header_style)
    workbook.add_named_style(body_style)

    for sheet_name, rows in sheets.items():
        worksheet = workbook.create_sheet(title=sheet_name[:31])

        if not rows:
            worksheet.append(["Sem dados"])
            worksheet["A1"].style = "leitorbi_header"
            continue

        if sheet_name == "Resumo":
            headers = ["Indicador", "Valor"]
            worksheet.append(headers)
            widths = [len(item) for item in headers]
            for key, value in rows[0].items():
                values = [key, value]
                worksheet.append(values)
                widths = [max(widths[index], min(len(text(value)), 58)) for index, value in enumerate(values)]
        else:
            headers = list(rows[0].keys())
            worksheet.append(headers)
            widths = [len(item) for item in headers]
            for row in rows:
                values = [row.get(header, "") for header in headers]
                worksheet.append(values)
                widths = [max(widths[index], min(len(text(value)), 58)) for index, value in enumerate(values)]

        for cell in worksheet[1]:
            cell.style = "leitorbi_header"
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.style = "leitorbi_body"

        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 60)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
